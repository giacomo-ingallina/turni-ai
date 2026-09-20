# =====================================================================
#  INDISPONIBILITÀ — crea il file da far compilare alle persone
#  Da incollare in una cella di Google Colab e avviare con il tasto ▶
#  Modifica solo la parte "DATI DA COMPILARE" qui sotto.
# =====================================================================

# ----------------------- DATI DA COMPILARE ---------------------------
# Si scrivono a parole, come nelle schede.

MESE_E_ANNO_DI_PARTENZA = """ottobre 2027"""
QUANTI_MESI = 12

NOMI_DELLE_PERSONE = """
ANCONA F, BARKI, BIONDI, BOGNONI, COLOMBI, CONVERSANO, CUNSOLO, FIORE,
INGALLINA, MARGONATO, MUSCI, PACI, STELLA
"""

# CODICI PER UNA SINGOLA ATTIVITÀ — facoltativo
#
# I codici di indisponibilità sono cinque e non si cambiano:
#   X  = non disponibile tutto il giorno
#   Xm = non disponibile la mattina
#   Xp = non disponibile il pomeriggio
#   Xg = non disponibile di giorno, cioè mattina e pomeriggio: la notte no
#   Xn = non disponibile la notte
#
# Qui si aggiungono codici di un altro tipo: non bloccano un orario, ma UNA
# SOLA ATTIVITÀ, nel giorno in cui vengono scritti. Chi li usa resta
# disponibile per tutto il resto di quella giornata.
# Uno per riga, nella forma  sigla = descrizione:
#
#   noG = indisponibile per il turno GUARDIA quel giorno
#
# IMPORTANTE: dentro la descrizione deve comparire il nome di un'attività
# dichiarata in «Crea il file Excel con i turni vuoti», scritto come lì
# (GUARDIA, non GUARDIE né «la guardia»). È da quel nome che si capisce che
# cosa bloccare: se non corrisponde a nessuna attività il codice non funziona,
# e lo script dell'orario si ferma e te lo dice.
#
# Scrivi «nessuno» se non ti servono.
CODICI_IN_PIU_TESTO = """
noG = indisponibile per il turno GUARDIA quel giorno
"""


# --------------------- FINE DATI DA COMPILARE ------------------------


# ------------- lettura di quello che hai scritto qui sopra ------------
import re, unicodedata

MESI = {"gennaio":1,"febbraio":2,"marzo":3,"aprile":4,"maggio":5,"giugno":6,
        "luglio":7,"agosto":8,"settembre":9,"ottobre":10,"novembre":11,"dicembre":12}

class Problema(Exception):
    pass

def _pulisci(t):
    t = unicodedata.normalize("NFC", str(t))
    return t.replace("\u2014", "-").replace("\u2013", "-").replace("\u2019", "'")

# ---------------------------------------------------------------- mese e anno
def leggi_mese(testo):
    t = _pulisci(testo).strip().lower()
    m = re.search(r"([a-zàèéìòù]+)\D*(\d{4})", t)
    if not m or m.group(1) not in MESI:
        raise Problema(f"non capisco il mese e l'anno da «{testo.strip()}». "
                       f"Scrivili così: novembre 2026")
    return MESI[m.group(1)], int(m.group(2)), m.group(1)

# ---------------------------------------------------------------- persone
def leggi_persone(testo):
    nomi = [n.strip() for n in re.split(r"[,;\n]", _pulisci(testo)) if n.strip()]
    nomi = [n.upper() for n in nomi]
    if not nomi:
        raise Problema("non ho trovato nessun nome")
    doppi = {n for n in nomi if nomi.count(n) > 1}
    if doppi:
        raise Problema("questi nomi compaiono due volte: " + ", ".join(sorted(doppi)))
    return nomi

# ---------------------------------------------------------------- codici
# I cinque codici di base sono fissi: non dipendono da come è fatto l'orario
# e non si possono cambiare. I codici in più bloccano una attività, mai una
# parte della giornata.
BASE = [("X",  "non disponibile tutto il giorno"),
        ("Xm", "non disponibile la mattina"),
        ("Xp", "non disponibile il pomeriggio"),
        ("Xg", "non disponibile di giorno, cioè mattina e pomeriggio: "
               "la notte resta disponibile"),
        ("Xn", "non disponibile la notte")]

def leggi_codici(testo):
    t = _pulisci(testo).strip()
    if not t or t.lower() in ("nessuno", "nessun", "no", "-"):
        return [], []
    voci, corrente = [], None
    for riga in t.split("\n"):
        m = re.match(r"\s*([A-Za-z][A-Za-z0-9]{0,4})\s*[=:]\s*(.*)$", riga)
        if m:
            if corrente: voci.append(corrente)
            corrente = [m.group(1), m.group(2).strip()]
        elif corrente:
            corrente[1] += " " + riga.strip()
    if corrente: voci.append(corrente)
    if not voci:
        raise Problema("non riconosco nessun codice. Scrivili così, uno per riga: "
                       "noG = indisponibile per il turno GUARDIA quel giorno")
    fuori, avvisi = [], []
    for sigla, descr in voci:
        d = re.sub(r"\s+", " ", descr).strip().rstrip(".")
        if not d:
            raise Problema(f"il codice «{sigla}» non ha una descrizione. Scrivi che "
                           f"cosa blocca: «{sigla} = indisponibile per il turno "
                           f"GUARDIA quel giorno».")
        if re.search(r"\b(mattin|pomerigg|nott|tutto il giorno|intera giornata|"
                     r"giorno success|giorno dopo)", d.lower()):
            raise Problema(
                f"il codice «{sigla}» sembra bloccare una parte della giornata. Per "
                f"quelle ci sono già i codici fissi X, Xm, Xp, Xg, Xn, che non si "
                f"cambiano: i codici in più bloccano UNA ATTIVITÀ, non un orario. "
                f"Una notte svolta altrove si segna con due X, sul giorno in cui "
                f"comincia e su quello dopo.")
        fuori.append((sigla, d))
    sigle = [s for s, _ in fuori]
    doppie = {s for s in sigle if sigle.count(s) > 1}
    if doppie:
        raise Problema("questi codici compaiono due volte: " + ", ".join(sorted(doppie)))
    scontro = set(sigle) & {c for c, _ in BASE}
    if scontro:
        raise Problema("questi codici esistono già di base e non vanno ridichiarati: "
                       + ", ".join(sorted(scontro)))
    return fuori, avvisi


try:
    _nm, ANNO, _nome_mese = leggi_mese(MESE_E_ANNO_DI_PARTENZA)
    MESE = _nome_mese.upper()
    PERSONE = leggi_persone(NOMI_DELLE_PERSONE)
    CODICI_IN_PIU, _av = leggi_codici(CODICI_IN_PIU_TESTO)
    NOME_FILE = f"Indisponibilita_da_{_nome_mese.capitalize()}_{ANNO}.xlsx"
except Problema as _e:
    print("NON POSSO PROCEDERE:", _e)
    raise SystemExit
for _a in _av:
    print("Avviso:", _a)

import calendar
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NOMI = ["", "Gennaio","Febbraio","Marzo","Aprile","Maggio","Giugno",
        "Luglio","Agosto","Settembre","Ottobre","Novembre","Dicembre"]
GG = ["lunedì","martedì","mercoledì","giovedì","venerdì","sabato","domenica"]

LEGENDA = BASE + CODICI_IN_PIU
CODICI = [c for c, _ in LEGENDA]

ISTRUZIONI = [
    "1. Cerca il foglio del mese giusto (in basso) e la colonna con il tuo nome.",
    "2. Sulla riga del giorno, scegli il codice dal menu a tendina.",
    "3. I giorni in cui sei disponibile vanno lasciati VUOTI.",
    "",
    "Esempi:",
    "  il giorno 3 sei assente solo la mattina: Xm sulla riga del 3",
    "  il giorno 10 sei assente tutto il giorno: X sulla riga del 10",
    "  il giorno 12 sei assente di giorno ma la notte puoi: Xg sulla riga del 12",
    "",
    "Una notte svolta altrove occupa due giornate: si segna con due X,",
    "una sul giorno in cui comincia e una sul giorno dopo, quando smonti.",
]
if CODICI_IN_PIU:
    ISTRUZIONI += ["",
        "I codici " + ", ".join(c for c, _ in CODICI_IN_PIU) + " sono diversi dagli",
        "altri: non bloccano una parte della giornata, ma una sola attività in",
        "quel giorno. Chi li usa resta disponibile per tutto il resto."]

NOTE_PREMESSA = [
    "Spazio libero, facoltativo: si può lasciare vuoto.",
    "Serve a chi prepara l'orario, per le ultime correzioni a mano.",
    "Non viene incollato nel file dei turni e non lo legge nessun programma:",
    "quello che conta per l'assegnazione sono i codici, non queste righe.",
    "Il file gira tra i colleghi, quindi lo leggono tutti: scrivi solo",
    "quello che ti va di far sapere, e non serve dire il motivo di un'assenza.",
    "Esempi: «il 10 ho messo Xm ma riesco ad arrivare per le 11»,",
    "«il 14 ho un impegno leggero, se serve posso dare una mano».",
]

# Excel ammette al massimo 255 caratteri nel messaggio di aiuto della tendina:
# oltre quel limite considera il file danneggiato e lo ripara all'apertura.
AIUTO = " | ".join(f"{c} = {s}" for c, s in LEGENDA)
if len(AIUTO) > 250:
    AIUTO = ("Codici ammessi: " + ", ".join(CODICI) +
             ". Il significato di ciascuno è nella legenda, a destra del foglio.")

CAL = "Calibri"
f_n = Font(name=CAL, size=11); f_b = Font(name=CAL, size=11, bold=True)
f_p = Font(name=CAL, size=10)
f_i = Font(name=CAL, size=10, italic=True)
we   = PatternFill("solid", fgColor="DDEBF7")
gr   = PatternFill("solid", fgColor="D9D9D9")
comp = PatternFill("solid", fgColor="FFF2CC")
mid  = Alignment(horizontal="center", vertical="center", wrap_text=False)
sx   = Alignment(horizontal="left",   vertical="center", wrap_text=False)
bordo = Border(*[Side(style="thin", color="BFBFBF")] * 4)

wb = Workbook(); wb.remove(wb.active)
P0 = 3
P1 = P0 + len(PERSONE) - 1
C_LEG = P1 + 2
LP0, LP1 = get_column_letter(P0), get_column_letter(P1)

anno, mese = ANNO, _nm
for _ in range(QUANTI_MESI):
    ng = calendar.monthrange(anno, mese)[1]
    R0, R1 = 4, 4 + ng - 1
    ws = wb.create_sheet(f"{NOMI[mese]} {anno}")
    ws["A1"], ws["B1"], ws["A2"] = NOMI[mese].upper(), anno, mese
    ws.cell(row=2, column=P0, value="INDISPONIBILITÀ")
    ws.merge_cells(start_row=2, start_column=P0, end_row=2, end_column=P1)
    ws.cell(row=3, column=1, value="GIORNO"); ws.cell(row=3, column=2, value="DATA")
    for i, p in enumerate(PERSONE):
        ws.cell(row=3, column=P0 + i, value=p)
    for i in range(ng):
        r = R0 + i
        ws.cell(row=r, column=1, value='=CHOOSE(WEEKDAY(DATE($B$1,$A$2,B%d),2),"%s")'
                % (r, '","'.join(GG)))
        ws.cell(row=r, column=2, value=i + 1)

    for r in (1, 2, 3):
        for c in range(1, P1 + 1):
            cel = ws.cell(row=r, column=c); cel.font = f_b; cel.alignment = mid
            if r == 3: cel.fill = gr; cel.border = bordo
    for i in range(ng):
        r = R0 + i; fine = calendar.weekday(anno, mese, i + 1) >= 5
        for c in range(1, P1 + 1):
            cel = ws.cell(row=r, column=c)
            cel.border = bordo; cel.alignment = mid; cel.font = f_n
            cel.fill = gr if c <= 2 else (we if fine else comp)
            if fine and c <= 2: cel.fill = we

    # ------------------------------------------------ legenda e istruzioni
    ws.cell(row=2, column=C_LEG, value="LEGENDA DEI CODICI").font = f_b
    ws.merge_cells(start_row=2, start_column=C_LEG, end_row=2, end_column=C_LEG + 1)
    ws.cell(row=3, column=C_LEG, value="CODICE").font = f_b
    ws.cell(row=3, column=C_LEG + 1, value="SIGNIFICATO").font = f_b
    for c in (C_LEG, C_LEG + 1):
        ws.cell(row=3, column=c).fill = gr; ws.cell(row=3, column=c).border = bordo
    for i, (c, sig) in enumerate(LEGENDA):
        a = ws.cell(row=4 + i, column=C_LEG, value=c); a.font = f_b
        a.alignment = mid; a.border = bordo
        b = ws.cell(row=4 + i, column=C_LEG + 1, value=sig); b.font = f_n
        b.alignment = sx; b.border = bordo
    r_ist = 4 + len(LEGENDA) + 2
    ws.cell(row=r_ist, column=C_LEG, value="COME COMPILARE").font = f_b
    for i, riga in enumerate(ISTRUZIONI):
        ws.cell(row=r_ist + 1 + i, column=C_LEG, value=riga).font = f_p

    # ------------------------------------------------------- blocco note
    r_note = r_ist + len(ISTRUZIONI) + 3
    ws.cell(row=r_note, column=C_LEG, value="NOTE — facoltative").font = f_b
    for i, riga in enumerate(NOTE_PREMESSA):
        ws.cell(row=r_note + 1 + i, column=C_LEG, value=riga).font = f_i
    r_tab = r_note + len(NOTE_PREMESSA) + 2
    ws.cell(row=r_tab, column=C_LEG, value="PERSONA").font = f_b
    ws.cell(row=r_tab, column=C_LEG + 1, value="NOTA").font = f_b
    for c in (C_LEG, C_LEG + 1):
        ws.cell(row=r_tab, column=c).fill = gr; ws.cell(row=r_tab, column=c).border = bordo
    N0, N1 = r_tab + 1, r_tab + len(PERSONE)
    for i, p in enumerate(PERSONE):
        a = ws.cell(row=N0 + i, column=C_LEG, value=p)
        a.font = f_b; a.alignment = sx; a.border = bordo; a.fill = gr
        b = ws.cell(row=N0 + i, column=C_LEG + 1)
        b.font = f_n; b.alignment = sx; b.border = bordo; b.fill = comp
    R_FONDO = N1

    def largh(testi, extra=2, minimo=6):
        return max(minimo, max((len(str(t)) for t in testi), default=0) + extra)
    ws.column_dimensions["A"].width = largh(GG)
    ws.column_dimensions["B"].width = 6
    for i, p in enumerate(PERSONE):
        ws.column_dimensions[get_column_letter(P0 + i)].width = largh([p])
    ws.column_dimensions[get_column_letter(C_LEG)].width = largh(
        ["CODICE", "PERSONA"] + CODICI + list(PERSONE) + ISTRUZIONI + NOTE_PREMESSA)
    ws.column_dimensions[get_column_letter(C_LEG + 1)].width = max(
        40, largh([s for _, s in LEGENDA]))
    ws.freeze_panes = "C4"

    dv = DataValidation(type="list", formula1='"' + ",".join(CODICI) + '"',
                        allow_blank=True, showErrorMessage=True, showInputMessage=True)
    dv.errorTitle = "Codice non ammesso"
    dv.error = "Codice non ammesso: seleziona un valore dall'elenco"
    dv.promptTitle = "Codici ammessi"
    dv.prompt = AIUTO
    ws.add_data_validation(dv); dv.add(f"{LP0}{R0}:{LP1}{R1}")

    # protezione: si blocca tutto tranne le due aree da compilare
    for row in ws.iter_rows(min_row=1, max_row=R_FONDO + 2,
                            min_col=1, max_col=C_LEG + 1):
        for cel in row:
            cel.protection = Protection(locked=True)
    for r in range(R0, R1 + 1):
        for c in range(P0, P1 + 1):
            ws.cell(row=r, column=c).protection = Protection(locked=False)
    for r in range(N0, N1 + 1):
        ws.cell(row=r, column=C_LEG + 1).protection = Protection(locked=False)
    ws.protection.sheet = True
    # le quattro azioni che devono restare PERMESSE (False = non vietata)
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    mese += 1
    if mese > 12: mese, anno = 1, anno + 1

wb.save(NOME_FILE)

# ---------------- pulizia di una parte che fa "riparare" il file --------------
import zipfile, os
def _togli_custom_props(percorso):
    tmp = percorso + ".tmp"
    with zipfile.ZipFile(percorso) as z:
        dati = {n: z.read(n) for n in z.namelist()}
    if "docProps/custom.xml" in dati:
        del dati["docProps/custom.xml"]
        dati["_rels/.rels"] = re.sub(
            r'<Relationship[^>]*custom-properties[^>]*/>', '',
            dati["_rels/.rels"].decode()).encode()
        dati["[Content_Types].xml"] = re.sub(
            r'<Override[^>]*docProps/custom\.xml[^>]*/>', '',
            dati["[Content_Types].xml"].decode()).encode()
    for nome in list(dati):
        if not nome.startswith("xl/worksheets/sheet"):
            continue
        testo = dati[nome].decode()
        m = re.search(r'<c r="A2"[^>]*><v>(\d+)</v>', testo)
        m2 = re.search(r'<c r="B1"[^>]*><v>(\d+)</v>', testo)
        if m and m2:
            mm, aa = int(m.group(1)), int(m2.group(1))
            for i in range(calendar.monthrange(aa, mm)[1]):
                r = 4 + i
                g = GG[calendar.weekday(aa, mm, i + 1)]
                testo = re.sub(
                    r'(<c r="A%d"[^>]*?)(\s+t="[^"]*")?(>.*?</f>)<v></v>' % r,
                    lambda x, g=g: x.group(1) + ' t="str"' + x.group(3) + "<v>" + g + "</v>",
                    testo, count=1, flags=re.S)
        dati[nome] = testo.encode()
    wbx = dati["xl/workbook.xml"].decode()
    wbx = re.sub(r"<calcPr[^>]*/>", '<calcPr calcId="191029" calcMode="auto" '
                 'fullCalcOnLoad="0" forceFullCalc="0"/>', wbx)
    if "<calcPr" not in wbx:
        wbx = wbx.replace("</workbook>", '<calcPr calcId="191029" calcMode="auto"/>'
                          "</workbook>")
    dati["xl/workbook.xml"] = wbx.encode()
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in dati.items():
            z.writestr(n, d)
    os.replace(tmp, percorso)
_togli_custom_props(NOME_FILE)

# ------------------------- VERIFICHE ---------------------------------
v = load_workbook(NOME_FILE)
def esito(t, ok, extra=""):
    print(("  OK   " if ok else "  ERRORE ") + t + ((" — " + extra) if extra else ""))

print("=== VERIFICHE SUL FILE SALVATO ===")
print("  file:", NOME_FILE)
esito(f"fogli creati: {len(v.sheetnames)}", len(v.sheetnames) == QUANTI_MESI,
      v.sheetnames[0] + " … " + v.sheetnames[-1])
righe_ok, primo_ok, nomi_ok, prot_ok, opz_ok, dv_ok, pwd_ok = ([] for _ in range(7))
note_ok = []
rif = None
for nome in v.sheetnames:
    w = v[nome]; a, m = w["B1"].value, w["A2"].value
    ng = calendar.monthrange(a, m)[1]
    r = 4
    while w.cell(row=r, column=2).value is not None: r += 1
    righe_ok.append(r - 4 == ng)
    primo_ok.append(str(w["A4"].value).startswith("=CHOOSE"))
    n = [w.cell(row=3, column=c).value for c in range(P0, P1 + 1)]
    if rif is None: rif = n
    nomi_ok.append(n == rif)
    p = w.protection
    prot_ok.append(p.sheet); pwd_ok.append(p.password is None)
    opz_ok.append(not p.selectLockedCells and not p.selectUnlockedCells
                  and not p.formatColumns and not p.formatRows)
    d = list(w.data_validations.dataValidation)
    dv_ok.append(len(d) == 1 and d[0].formula1 == '"' + ",".join(CODICI) + '"')
    # blocco note: una riga per persona, cella della nota sbloccata
    trovati, sbloccate = [], True
    for rr in range(1, w.max_row + 1):
        if w.cell(row=rr, column=C_LEG).value == "PERSONA":
            for i2 in range(len(PERSONE)):
                trovati.append(w.cell(row=rr + 1 + i2, column=C_LEG).value)
                if w.cell(row=rr + 1 + i2, column=C_LEG + 1).protection.locked:
                    sbloccate = False
            break
    note_ok.append(trovati == PERSONE and sbloccate)
esito("tutti i mesi hanno il numero di giorni giusto", all(righe_ok))
esito("giorno della settimana calcolato con formula in tutti i fogli", all(primo_ok))
esito("nomi delle persone identici e nella stessa posizione", all(nomi_ok))
esito(f"tendina con i soli codici {', '.join(CODICI)}", all(dv_ok))
esito("blocco note: una riga per persona, spazio sbloccato", all(note_ok))
esito("protezione attiva su tutti i fogli", all(prot_ok))
esito("nessuna password su nessun foglio", all(pwd_ok))
esito("permesso selezionare, compilare, allargare colonne e righe", all(opz_ok))
w0 = v[v.sheetnames[0]]
esito("celle da compilare sbloccate",
      not w0.cell(row=4, column=P0).protection.locked)
esito("celle di giorno e data bloccate", w0["A4"].protection.locked)
esito("nessuna attività o turno nel file",
      all(w0.cell(row=3, column=c).value in (None, "GIORNO", "DATA") or
          w0.cell(row=3, column=c).value in PERSONE or c >= C_LEG
          for c in range(1, C_LEG + 2)))
d0 = list(w0.data_validations.dataValidation)[0]
esito(f"messaggio di aiuto entro i 255 caratteri di Excel: {len(d0.prompt or '')}",
      len(d0.prompt or "") <= 255)
esito(f"messaggio di errore entro i 255 caratteri: {len(d0.error or '')}",
      len(d0.error or "") <= 255)
esito("nessun testo a capo",
      not any(c.alignment.wrap_text for row in w0.iter_rows() for c in row))

print()
print("APRI IL FILE E PROVA QUESTO: clicca su una cella sotto un nome e scrivi,")
print("poi prova a scrivere una nota nella colonna NOTA, in basso a destra.")

try:
    from google.colab import files
    files.download(NOME_FILE)
except Exception:
    print("\n(fuori da Colab: il file è stato salvato come", NOME_FILE + ")")
