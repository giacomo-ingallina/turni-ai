# =====================================================================
#  ORARIO DEI TURNI — crea il file Excel del mese, vuoto
#  Da incollare in una cella di Google Colab e avviare con il tasto ▶
#  Modifica solo la parte "DATI DA COMPILARE" qui sotto.
# =====================================================================

# ----------------------- DATI DA COMPILARE ---------------------------
# Si scrivono a parole, come nelle schede. Modifica solo il testo fra le
# virgolette triple: il programma lo legge e ti dice se qualcosa non torna.

MESE_E_ANNO = """novembre 2026"""

ATTIVITA_DA_COPRIRE = """
STANZA 23 - 2 turni: MAT, POM
STANZA 24 - 2 turni: MAT, POM
STANZA 25 - 2 turni: MAT, POM
STANZA 26 - 2 turni: MAT, POM
STANZA 43 - 1 turno: MAT
STANZA 31 - 1 turno: MAT
AMB ESTERNO - 1 turno: POM
GUARDIA - 3 turni: MAT, POM, NOTTE
"""

NOMI_DELLE_PERSONE = """
LEONI, ROSSI M, ROSSI L, GALLINA, RUSSO, VERDI, HU,
MARINI, CONTI A, CONTI S, FERRARI, GRECO, SANNA, LI
"""

# CODICI PER UNA SINGOLA ATTIVITÀ — facoltativo.
# Qui questi codici servono a due cose sole: finire nella legenda del file e
# far controllare che il nome dell'attività sia scritto giusto. Non entrano in
# nessuna formula e non rendono nessuno indisponibile: quel vincolo lo applica
# il chatbot in fase di assegnazione, leggendo la scheda del PROMPT C.
# I cinque codici di indisponibilità sono fissi (X, Xm, Xp, Xg, Xn) e non si
# cambiano. Dentro la descrizione deve comparire il nome di un'attività
# dichiarata qui sopra, scritto come lì. «nessuno» se non ne servono.
CODICI_IN_PIU_TESTO = """
noG = indisponibile per il turno GUARDIA quel giorno
"""


# --------------------- FINE DATI DA COMPILARE ------------------------


# ------------- lettura di quello che hai scritto qui sopra ------------
import re, unicodedata

MESI = {"gennaio":1,"febbraio":2,"marzo":3,"aprile":4,"maggio":5,"giugno":6,
        "luglio":7,"agosto":8,"settembre":9,"ottobre":10,"novembre":11,"dicembre":12}

class Problema(Exception):
    pass

def _pulisci(t):
    t = unicodedata.normalize("NFC", str(t))
    return t.replace("\u2014", "-").replace("\u2013", "-").replace("\u2019", "'")

# ---------------------------------------------------------------- mese e anno
def leggi_mese(testo):
    t = _pulisci(testo).strip().lower()
    m = re.search(r"([a-zàèéìòù]+)\D*(\d{4})", t)
    if not m or m.group(1) not in MESI:
        raise Problema(f"non capisco il mese e l'anno da «{testo.strip()}». "
                       f"Scrivili così: novembre 2026")
    return MESI[m.group(1)], int(m.group(2)), m.group(1)

# ---------------------------------------------------------------- attività
GIORNI_SETT = {"lunedì": 0, "lunedi": 0, "martedì": 1, "martedi": 1,
               "mercoledì": 2, "mercoledi": 2, "giovedì": 3, "giovedi": 3,
               "venerdì": 4, "venerdi": 4, "sabato": 5, "domenica": 6}

# «GIORNO solo sabato e domenica»  -> vale per tutte le attività
SOLO = re.compile(r"^\s*([A-Za-z0-9 ,ed]+?)\s+solo\s+(.+?)\s*\.?\s*$", re.I)
# «GUARDIA, turno GIORNO: solo sabato e domenica»  -> vale per quella sola attività
SOLO_ATT = re.compile(r"^\s*(.+?)\s*,\s*turn[oi]\s+(.+?)\s*:\s*solo\s+(.+?)\s*\.?\s*$",
                      re.I)

def leggi_giorni(testo):
    """«sabato e domenica», «dal lunedì al venerdì», «nei giorni feriali»,
    «nel weekend»: restituisce l'insieme dei giorni, 0 = lunedì."""
    t = testo.lower()
    if re.search(r"\bferial", t):
        return {0, 1, 2, 3, 4}
    if re.search(r"\bweekend|fine settimana|festiv", t):
        return {5, 6}
    m = re.search(r"\bdal\s+([a-zì]+)\s+al\s+([a-zì]+)", t)
    if m and m.group(1) in GIORNI_SETT and m.group(2) in GIORNI_SETT:
        a, b = GIORNI_SETT[m.group(1)], GIORNI_SETT[m.group(2)]
        return set(range(a, b + 1)) if a <= b else set(range(a, 7)) | set(range(0, b + 1))
    trovati = {GIORNI_SETT[g] for g in GIORNI_SETT if re.search(r"\b" + g + r"\b", t)}
    return trovati or None

OCCUPA = re.compile(
    r"^\s*([A-Za-z0-9 ]+?)\s+occupa\s+"
    r"(tutta la giornata|tutto il giorno|la giornata e quella dopo|"
    r"il giorno e quello dopo)\s*\.?\s*$", re.I)

def leggi_attivita(testo):
    """Una attività per riga: NOME - 2 turni: MAT, POM
    Una riga può invece dire in quali giorni della settimana esiste un turno:
    «MAT e POM solo dal lunedì al venerdì», oppure, per una sola attività,
    «GUARDIA, turno GIORNO: solo sabato e domenica»."""
    fuori, avvisi = [], []
    esclusivi, prolungati = [], []
    giorni, giorni_att = {}, {}
    for riga in _pulisci(testo).split("\n"):
        r = riga.strip().strip("-•").strip()
        if not r:
            continue
        m = SOLO_ATT.match(r)
        if m:
            gg = leggi_giorni(m.group(3))
            if gg is None:
                raise Problema(
                    f"nella riga «{r}» non capisco in quali giorni. Scrivi per esempio "
                    f"«GUARDIA, turno GIORNO: solo sabato e domenica».")
            att = m.group(1).strip().upper()
            for turno in re.split(r"[,;]| e ", m.group(2)):
                turno = turno.strip().upper()
                if turno:
                    giorni_att[(att, turno)] = gg
            continue
        m = SOLO.match(r)
        if m and " solo " in r.lower() and ":" not in r:
            gg = leggi_giorni(m.group(2))
            if gg is None:
                raise Problema(
                    f"nella riga «{r}» non capisco in quali giorni. Scrivi per esempio "
                    f"«GIORNO solo sabato e domenica», «MAT e POM solo dal lunedì al "
                    f"venerdì» oppure «NOTTE solo nel weekend».")
            for turno in re.split(r"[,;]| e ", m.group(1)):
                turno = turno.strip().upper()
                if turno:
                    giorni[turno] = gg
            continue
        m = OCCUPA.match(r)
        if m:
            turno = m.group(1).strip().upper()
            esclusivi.append(turno)
            if "dopo" in m.group(2).lower():
                prolungati.append(turno)
            continue
        # tutto quello che sta dopo i due punti sono le etichette dei turni
        if ":" not in r:
            raise Problema(f"nella riga «{r}» non trovo i due punti. "
                           f"Scrivila così: STANZA 23 - 2 turni: MAT, POM")
        testa, coda = r.split(":", 1)
        turni = [x.strip().upper() for x in re.split(r"[,;/]| e ", coda) if x.strip()]
        if not turni:
            raise Problema(f"nella riga «{r}» non trovo i nomi dei turni")
        # il nome è quello che precede il trattino, o l'indicazione «N turni»
        nome = re.split(r"\s*-\s*|\s+\d+\s*turn", testa)[0].strip()
        if not nome:
            raise Problema(f"nella riga «{r}» non trovo il nome dell'attività")
        dichiarati = re.search(r"(\d+)\s*turn", testa)
        if dichiarati and int(dichiarati.group(1)) != len(turni):
            avvisi.append(f"«{nome}»: dichiari {dichiarati.group(1)} turni ma ne elenchi "
                          f"{len(turni)} ({', '.join(turni)}). Uso quelli elencati.")
        fuori.append((nome.upper(), turni))
    if not fuori:
        raise Problema("non ho trovato nessuna attività")
    tutti = {t for _, ts in fuori for t in ts}
    ignoti = [t for t in esclusivi if t not in tutti]
    if ignoti:
        raise Problema("dichiari che " + ", ".join(ignoti) + " occupa tutta la giornata, "
                       "ma quel turno non compare in nessuna attività. Controlla come "
                       "l'hai scritto: deve essere identico all'etichetta del turno.")
    ignoti = [t for t in giorni if t not in tutti]
    if ignoti:
        raise Problema("dichiari i giorni di " + ", ".join(ignoti) + ", ma quel turno non "
                       "compare in nessuna attività. Il nome deve essere identico "
                       "all'etichetta del turno.")
    coppie = {(n, t) for n, ts in fuori for t in ts}
    sbagliate = [f"{a}, turno {t}" for (a, t) in giorni_att if (a, t) not in coppie]
    if sbagliate:
        nomi = ", ".join(n for n, _ in fuori)
        raise Problema("queste dichiarazioni non corrispondono a nessuna attività con quel "
                       "turno: " + "; ".join(sbagliate) + ". Le attività sono: " + nomi +
                       ". Nome e turno devono essere scritti come nella riga dell'attività.")
    return fuori, esclusivi, prolungati, giorni, giorni_att, avvisi

# ---------------------------------------------------------------- persone
def leggi_persone(testo):
    nomi = [n.strip() for n in re.split(r"[,;\n]", _pulisci(testo)) if n.strip()]
    nomi = [n.upper() for n in nomi]
    if not nomi:
        raise Problema("non ho trovato nessun nome")
    doppi = {n for n in nomi if nomi.count(n) > 1}
    if doppi:
        raise Problema("questi nomi compaiono due volte: " + ", ".join(sorted(doppi)))
    return nomi

# ---------------------------------------------------------------- codici
def leggi_codici(testo, attivita=None):
    """I codici in più bloccano UNA ATTIVITA' nel giorno in cui sono scritti.
    Non bloccano parti della giornata: per quelle ci sono X, Xm, Xp, Xg, Xn."""
    t = _pulisci(testo).strip()
    if not t or t.lower() in ("nessuno", "nessun", "no", "-"):
        return [], []
    voci, corrente = [], None
    for riga in t.split("\n"):
        m = re.match(r"\s*([A-Za-z][A-Za-z0-9]{0,4})\s*[=:]\s*(.*)$", riga)
        if m:
            if corrente: voci.append(corrente)
            corrente = [m.group(1), m.group(2).strip()]
        elif corrente:
            corrente[1] += " " + riga.strip()
    if corrente: voci.append(corrente)
    if not voci:
        raise Problema("non riconosco nessun codice. Scrivili così, uno per riga: "
                       "noG = indisponibile per il turno GUARDIA quel giorno")
    nomi_att = [n for n, _ in (attivita or [])]
    fuori, avvisi = [], []
    for sigla, descr in voci:
        d = re.sub(r"\s+", " ", descr).strip().rstrip(".")
        trovata = None
        for n in sorted(nomi_att, key=len, reverse=True):
            if re.search(r"(?<![A-Za-z0-9])" + re.escape(n) + r"(?![A-Za-z0-9])",
                         d, re.IGNORECASE):
                trovata = n
                break
        if trovata is None:
            raise Problema(
                f"del codice «{sigla}» non capisco quale attività blocca. Nella "
                f"descrizione deve comparire il nome di una delle attività, scritto "
                f"identico: " + ", ".join(nomi_att) + ". Per esempio: "
                f"«{sigla} = indisponibile per il turno GUARDIA quel giorno».")
        if re.search(r"\bblocca\b.*\b(mattin|pomerigg|nott|tutto il giorno|"
                     r"giorno success|giorno dopo)", d.lower()):
            raise Problema(
                f"il codice «{sigla}» sembra voler bloccare una parte della giornata. "
                f"Per quelle ci sono già i codici fissi X, Xm, Xp, Xg, Xn, che non "
                f"si possono cambiare: i codici in più bloccano una attività, non un "
                f"orario. Una notte fatta altrove si segna con due X, sul giorno e su "
                f"quello dopo.")
        fuori.append((sigla, d, "turno:" + trovata))
    sigle = [s for s, _, _ in fuori]
    doppie = {s for s in sigle if sigle.count(s) > 1}
    if doppie:
        raise Problema("questi codici compaiono due volte: " + ", ".join(sorted(doppie)))
    scontro = set(sigle) & {"X", "Xm", "Xp", "Xg", "Xn"}
    if scontro:
        raise Problema("questi codici esistono già di base e non vanno ridichiarati: "
                       + ", ".join(sorted(scontro)))
    doppio_att = {}
    for sg, _, q in fuori:
        doppio_att.setdefault(q[6:], []).append(sg)
    for att, sg in doppio_att.items():
        if len(sg) > 1:
            avvisi.append(f"i codici {', '.join(sg)} bloccano la stessa attività "
                          f"({att}): controlla che sia voluto.")
    return fuori, avvisi

try:
    NMESE, ANNO, _nome_mese = leggi_mese(MESE_E_ANNO)
    MESE = _nome_mese.upper()
    ATTIVITA, _escl, _prol, _giorni, _giorni_att, _av1 = leggi_attivita(ATTIVITA_DA_COPRIRE)
    PERSONE = leggi_persone(NOMI_DELLE_PERSONE)
    _ammesse = ["MAT", "POM", "GIORNO", "NOTTE"]
    _ignoti = []
    for _n, _ts in ATTIVITA:
        for _t in _ts:
            if _t not in _ammesse and _t not in _ignoti: _ignoti.append(_t)
    if _ignoti:
        raise Problema(
            "questi turni non esistono: " + ", ".join(_ignoti) + ". I turni sono "
            "quattro e non si possono cambiare: MAT (resta libero il pomeriggio), "
            "POM (resta libero la mattina), GIORNO (occupa tutta la giornata), "
            "NOTTE (occupa quel giorno e anche quello dopo, perche' si smonta). "
            "Se un'attività è aperta solo in certi giorni, scrivilo con una riga "
            "a parte, per esempio «AMB ESTERNO, turno POM: solo dal lunedì al venerdì».")
    if _escl or _prol:
        raise Problema(
            "non serve più dichiarare che un turno occupa tutta la giornata: lo dice "
            "già il suo nome. Usa GIORNO per un turno che occupa l'intera giornata e "
            "NOTTE per uno che occupa anche il giorno dopo, e togli la riga «"
            + (_escl + _prol)[0] + " occupa ...».")
    CODICI_IN_PIU, _av2 = leggi_codici(CODICI_IN_PIU_TESTO, ATTIVITA)
    # il nome del file segue il mese: così cambiando il mese non si sovrascrive
    # per sbaglio l'orario precedente
    NOME_FILE = f"Turni_{_nome_mese.capitalize()}_{ANNO}.xlsx"
except Problema as _e:
    print("NON POSSO PROCEDERE:", _e)
    raise SystemExit
for _a in _av1 + _av2:
    print("Avviso:", _a)

import calendar, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NG = calendar.monthrange(ANNO, NMESE)[1]
R0, R1 = 4, 4 + NG - 1

# Parti della giornata: sono le etichette dei turni, nell'ordine in cui
# compaiono nelle attività. Sono quattro e sempre le stesse — MAT, POM,
# GIORNO, NOTTE — e a ognuna corrisponde un suffisso (_m, _p, _g, _n).
PARTI = []
for _, _ts in ATTIVITA:
    for _t in _ts:
        if _t not in PARTI: PARTI.append(_t)

PARTI_AMMESSE = ["MAT", "POM", "GIORNO", "NOTTE"]
SUF_PARTE = {"MAT": "_m", "POM": "_p", "GIORNO": "_g", "NOTTE": "_n"}
# I cinque codici sono fissi e non dipendono dai turni dichiarati.
BASE = [("X",  "non disponibile tutto il giorno", "giorno"),
        ("Xm", "non disponibile la mattina",      "parte:MAT"),
        ("Xp", "non disponibile il pomeriggio",   "parte:POM"),
        ("Xg", "non disponibile di giorno, cioè mattina e pomeriggio: "
               "la notte resta disponibile",      "parte:GIORNO"),
        ("Xn", "non disponibile la notte",        "parte:NOTTE")]
# Quali codici rendono impossibile ciascun turno. GIORNO copre mattina e
# pomeriggio, quindi lo bloccano Xm, Xp e Xg; Xg blocca anche MAT e POM.
BLOCCA = {"MAT":    ["Xm", "Xg"],
          "POM":    ["Xp", "Xg"],
          "GIORNO": ["Xm", "Xp", "Xg"],
          "NOTTE":  ["Xn"]}
LEGENDA = BASE + CODICI_IN_PIU
CODICI = [c for c, _, _ in LEGENDA]
INTERI = ["X"]          # blocca l'intera giornata
DOPO   = []             # nessun codice si estende al giorno dopo

# Le regole discendono dal nome del turno, senza doverle dichiarare:
#   MAT / POM ... mezze giornate, convivono fra loro
#   GIORNO ...... occupa l'intera giornata: chi lo fa non è libero per niente
#   NOTTE ....... occupa la giornata e anche quella dopo, perché si smonta
# i turni attivi solo in certi giorni della settimana: fuori da quei giorni la
# cella non va usata, e nessuno risulta libero per quel turno
GIORNI_PARTE = {p: _giorni.get(p) for p in PARTI}
GIORNI_ATT = dict(_giorni_att)

def attiva(parte, giorno_sett, attivita=None):
    """La regola dell'attività, se c'è, ha la precedenza su quella generale."""
    if attivita is not None and (attivita, parte) in GIORNI_ATT:
        return giorno_sett in GIORNI_ATT[(attivita, parte)]
    gg = GIORNI_PARTE.get(parte)
    return gg is None or giorno_sett in gg

def parti_attive(giorno_sett):
    """Una parte della giornata conta se almeno un'attività la usa quel giorno."""
    return [p for p in PARTI
            if any(attiva(p, giorno_sett, n) for n, ts in ATTIVITA if p in ts)]

ESCL = [p for p in PARTI if p in ("GIORNO", "NOTTE") or p in _escl]
DOPO_TURNO = [p for p in PARTI if p == "NOTTE" or p in _prol]

wb = Workbook(); ws = wb.active; ws.title = f"{MESE.capitalize()} {ANNO}"
CAL = "Calibri"
f_n = Font(name=CAL, size=11); f_b = Font(name=CAL, size=11, bold=True)
we   = PatternFill("solid", fgColor="DDEBF7")
gr   = PatternFill("solid", fgColor="D9D9D9")
corr = PatternFill("solid", fgColor="E2EFDA")
prec = PatternFill("solid", fgColor="FCE4D6")
app  = PatternFill("solid", fgColor="F2F2F2")
spento = PatternFill("solid", fgColor="BFBFBF")     # turno che quel giorno non c'è
mid  = Alignment(horizontal="center", vertical="center", wrap_text=False)
sx   = Alignment(horizontal="left",   vertical="center", wrap_text=False)
bordo = Border(*[Side(style="thin", color="BFBFBF")] * 4)

ws["A1"], ws["B1"], ws["A2"] = MESE, ANNO, NMESE

col = 3; mappa = {}
for nome, turni in ATTIVITA:
    a = col
    for t in turni:
        ws.cell(row=3, column=col, value=t); col += 1
    mappa[nome] = (a, col - 1)
    ws.cell(row=2, column=a, value=nome)
    if col - 1 > a:
        ws.merge_cells(start_row=2, start_column=a, end_row=2, end_column=col - 1)
T0, T1 = 3, col - 1
LT0, LT1 = get_column_letter(T0), get_column_letter(T1)

col += 1                                   # colonna vuota di separazione
C_CTRL, C_DISP, C_LIB = col, col + 1, col + 2
for c, t in ((C_CTRL, "CONTROLLO"), (C_DISP, "DISPONIBILI"), (C_LIB, "LIBERI")):
    ws.cell(row=3, column=c, value=t)
col += 3 + 4                               # più le quattro colonne vuote

C_GG, C_DD = col, col + 1
ws.cell(row=3, column=C_GG, value="GIORNO"); ws.cell(row=3, column=C_DD, value="DATA")
col += 2
P0 = col
for p in PERSONE:
    ws.cell(row=3, column=col, value=p); col += 1
P1 = col - 1
LP0, LP1 = get_column_letter(P0), get_column_letter(P1)
ws.cell(row=2, column=C_GG, value="GUARDIE E INDISPONIBILITÀ")
ws.merge_cells(start_row=2, start_column=C_GG, end_row=2, end_column=P1)
col += 1

C_LEG = col                                # legenda
ws.cell(row=2, column=C_LEG, value="LEGENDA CODICI").font = f_b
ws.merge_cells(start_row=2, start_column=C_LEG, end_row=2, end_column=C_LEG + 1)
ws.cell(row=3, column=C_LEG, value="CODICE"); ws.cell(row=3, column=C_LEG + 1, value="SIGNIFICATO")
for i, (c, sig, _) in enumerate(LEGENDA):
    ws.cell(row=4 + i, column=C_LEG, value=c).font = f_b
    ws.cell(row=4 + i, column=C_LEG + 1, value=sig).alignment = sx
col = C_LEG + 3

def blocco(c0, titolo, formule, riemp):
    ws.cell(row=3, column=c0, value="PERSONA")
    for i, (nome, _) in enumerate(ATTIVITA):
        ws.cell(row=3, column=c0 + 1 + i, value=nome)
    c1 = c0 + len(ATTIVITA) + 1
    ws.cell(row=3, column=c1, value="TOTALE")
    ws.cell(row=2, column=c0, value=titolo)
    ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c1)
    for i, p in enumerate(PERSONE):
        r = 4 + i
        ws.cell(row=r, column=c0, value=f"={get_column_letter(P0+i)}$3")
        if formule:
            rif = f"${get_column_letter(c0)}{r}"
            for j, (nome, _) in enumerate(ATTIVITA):
                a, b = mappa[nome]
                ws.cell(row=r, column=c0 + 1 + j, value=(
                    f"=COUNTIF(${get_column_letter(a)}${R0}:"
                    f"${get_column_letter(b)}${R1},{rif})"))
            ws.cell(row=r, column=c1, value=f"=COUNTIF(${LT0}${R0}:${LT1}${R1},{rif})")
    rt = 4 + len(PERSONE)
    ws.cell(row=rt, column=c0, value="TOTALE COMPLESSIVO")
    for c in range(c0 + 1, c1 + 1):
        L = get_column_letter(c)
        ws.cell(row=rt, column=c, value=f"=SUM({L}4:{L}{rt-1})")
    for r in range(2, rt + 1):
        for c in range(c0, c1 + 1):
            cel = ws.cell(row=r, column=c); cel.border = bordo; cel.alignment = mid
            cel.font = f_b if (r <= 3 or r == rt) else f_n
            cel.fill = gr if (r <= 3 or r == rt) else riemp
    return c0, c1, rt

CC0, CC1, RT = blocco(col, "TURNI MESE CORRENTE", True, corr)
PC0, PC1, _  = blocco(CC1 + 2, "TURNI MESE PRECEDENTE", False, prec)
ws.cell(row=RT + 2, column=PC0, value="Incollare qui SOLO VALORI (incolla speciale > "
        "valori) il blocco TURNI MESE CORRENTE del file del mese precedente"
        ).font = Font(name=CAL, size=10, italic=True)

A_D = PC1 + 2                              # aree di calcolo
A_L = A_D + len(PERSONE) + 1
for c0, tit in ((A_D, "AREA DI CALCOLO DISPONIBILI — non modificare"),
                (A_L, "AREA DI CALCOLO LIBERI — non modificare")):
    ws.cell(row=2, column=c0, value=tit).font = f_b
    ws.merge_cells(start_row=2, start_column=c0, end_row=2,
                   end_column=c0 + len(PERSONE) - 1)
    for i in range(len(PERSONE)):
        ws.cell(row=3, column=c0 + i,
                value=f"={get_column_letter(P0+i)}$3").font = f_b

def pezzo(i, r, con_turni):
    p = get_column_letter(P0 + i); cod = f"{p}{r}"; nome = f"{p}$3"
    # lo zero iniziale non è un vezzo: in Excel i valori logici valgono più di
    # qualsiasi numero, quindi FALSO>0 darebbe VERO. Sommando uno zero il
    # risultato è un numero, e il confronto torna a significare quello che dice.
    interi = "0+" + "+".join(f'({cod}="{c}")' for c in INTERI)
    if r > R0 and DOPO:
        interi += "+" + "+".join(f'({p}{r-1}="{c}")' for c in DOPO)
    def libero(et, riga=None):
        rr = r if riga is None else riga
        return f'(COUNTIFS(${LT0}$3:${LT1}$3,"{et}",${LT0}{rr}:${LT1}{rr},{nome})=0)'
    # chi ieri era di notte oggi smonta: non è libero per niente, esattamente
    # come per i codici prolungati
    smonto = ""
    if con_turni and r > R0:
        for e in DOPO_TURNO:
            smonto += f'*{libero(e, r - 1)}'
    attive = parti_attive(calendar.weekday(ANNO, NMESE, r - R0 + 1))
    pezzi = []
    for parte in attive:
        t = "*".join(f'({cod}<>"{x}")' for x in BLOCCA[parte])
        if con_turni:
            # non assegnato in questa parte, e in nessuna parte esclusiva;
            # se la parte è essa stessa esclusiva, in nessuna parte affatto
            da_escludere = attive if parte in ESCL else \
                           [parte] + [e for e in ESCL if e != parte and e in attive]
            for e in da_escludere:
                t += f'*{libero(e)}'
            t += smonto
        pezzi.append(t)
    tot = "0+" + "+".join(f"({x})" for x in pezzi)
    suff = "".join(f'&IF({x},"{SUF_PARTE[p]}","")'
                   for p, x in zip(attive, pezzi))
    return (f'=IF({interi}>0,"",'
            f'IF({tot}=0,"",'
            f'IF({tot}={len(pezzi)},{nome}&", ",'
            f'{nome}{suff}&", ")))')

def unione(c0, r):
    cat = "&".join(f"{get_column_letter(c0+i)}{r}" for i in range(len(PERSONE)))
    return f'=IF(LEN({cat})=0,"",LEFT({cat},LEN({cat})-2))'

GG = ["lunedì","martedì","mercoledì","giovedì","venerdì","sabato","domenica"]
for i in range(NG):
    r, g = R0 + i, i + 1
    ws.cell(row=r, column=1, value='=CHOOSE(WEEKDAY(DATE($B$1,$A$2,B%d),2),"%s")'
            % (r, '","'.join(GG)))
    ws.cell(row=r, column=2, value=g)
    ws.cell(row=r, column=C_CTRL, value=(
        f'=IF(SUMPRODUCT((${LT0}$3:${LT1}$3<>"")*(${LT0}{r}:${LT1}{r}<>"")*'
        f'(COUNTIFS(${LT0}{r}:${LT1}{r},${LT0}{r}:${LT1}{r},'
        f'${LT0}$3:${LT1}$3,${LT0}$3:${LT1}$3)>1))>0,"X","")'))
    ws.cell(row=r, column=C_GG, value=f"=A{r}")
    ws.cell(row=r, column=C_DD, value=f"=B{r}")
    for i2 in range(len(PERSONE)):
        ws.cell(row=r, column=A_D + i2, value=pezzo(i2, r, False))
        ws.cell(row=r, column=A_L + i2, value=pezzo(i2, r, True))
    ws.cell(row=r, column=C_DISP, value=unione(A_D, r))
    ws.cell(row=r, column=C_LIB,  value=unione(A_L, r))

ULT = A_L + len(PERSONE) - 1
for r in (1, 2, 3):
    for c in range(1, ULT + 1):
        cel = ws.cell(row=r, column=c); cel.font = f_b; cel.alignment = mid
        if r == 3: cel.fill = gr; cel.border = bordo
for i in range(NG):
    r = R0 + i; fine = calendar.weekday(ANNO, NMESE, i + 1) >= 5
    gs = calendar.weekday(ANNO, NMESE, i + 1)
    for c in list(range(1, C_LIB + 1)) + list(range(C_GG, P1 + 1)):
        cel = ws.cell(row=r, column=c); cel.border = bordo; cel.alignment = mid; cel.font = f_n
        if fine: cel.fill = we
    # i turni che quel giorno non esistono: cella barrata in grigio
    for nome_att, turni_att in ATTIVITA:
        a, b = mappa[nome_att]
        for k, t_ in enumerate(turni_att):
            if not attiva(t_, gs, nome_att):
                cel = ws.cell(row=r, column=a + k)
                cel.fill = spento
                cel.value = None
    for c in (1, 2, C_GG, C_DD):
        if not fine: ws.cell(row=r, column=c).fill = gr
    for c in range(A_D, ULT + 1):
        cel = ws.cell(row=r, column=c); cel.font = f_n; cel.alignment = mid; cel.fill = app

def largh(testi, extra=2, minimo=6):
    return max(minimo, max((len(str(t)) for t in testi if t is not None), default=0) + extra)
nome_max = largh(PERSONE)
ws.column_dimensions["A"].width = largh(GG)
ws.column_dimensions["B"].width = 6
for nome, turni in ATTIVITA:
    a, b = mappa[nome]
    per = max(nome_max, (len(nome) + 2) // max(1, b - a + 1))
    for c in range(a, b + 1):
        ws.column_dimensions[get_column_letter(c)].width = max(per, largh(turni))
ws.column_dimensions[get_column_letter(C_CTRL - 1)].width = 3
ws.column_dimensions[get_column_letter(C_CTRL)].width = largh(["CONTROLLO"])
elenco = ", ".join(PERSONE)
ws.column_dimensions[get_column_letter(C_DISP)].width = len(elenco) + 3
ws.column_dimensions[get_column_letter(C_LIB)].width = len(elenco) + 4 * len(PERSONE) + 3
ws.column_dimensions[get_column_letter(C_GG)].width = largh(GG)
ws.column_dimensions[get_column_letter(C_DD)].width = 6
for i, p in enumerate(PERSONE):
    ws.column_dimensions[get_column_letter(P0 + i)].width = largh([p])
ws.column_dimensions[get_column_letter(C_LEG)].width = largh(["CODICE"] + CODICI)
ws.column_dimensions[get_column_letter(C_LEG + 1)].width = largh([s for _, s, _ in LEGENDA])
for b0, b1 in ((CC0, CC1), (PC0, PC1)):
    ws.column_dimensions[get_column_letter(b0)].width = largh(["TOTALE COMPLESSIVO"] + PERSONE)
    for n2, (nome, _) in enumerate(ATTIVITA):
        ws.column_dimensions[get_column_letter(b0 + 1 + n2)].width = largh([nome])
    ws.column_dimensions[get_column_letter(b1)].width = largh(["TOTALE"])
for c in range(A_D, ULT + 1):
    ws.column_dimensions[get_column_letter(c)].width = nome_max + 2
    ws.column_dimensions[get_column_letter(c)].hidden = True

dv = DataValidation(type="list", formula1='"' + ",".join(CODICI) + '"',
                    allow_blank=True, showErrorMessage=True, showInputMessage=True)
dv.errorTitle = "Codice non ammesso"
dv.error = "Codice non ammesso: seleziona un valore dall'elenco"
dv.promptTitle = "Codici ammessi"
# Excel ammette al massimo 255 caratteri nel messaggio di aiuto della tendina:
# oltre quel limite considera il file danneggiato e lo ripara all'apertura.
aiuto = " | ".join(f"{c} = {s}" for c, s, _ in LEGENDA)
if len(aiuto) > 250:
    aiuto = "Codici ammessi: " + ", ".join(CODICI) + \
            ". Il significato di ciascuno è nella legenda, a destra del foglio."
dv.prompt = aiuto
ws.add_data_validation(dv); dv.add(f"{LP0}{R0}:{LP1}{R1}")

if any(GIORNI_PARTE.get(p) for p in PARTI) or GIORNI_ATT:
    _righe_nota = 4 + len(LEGENDA) + 1
    ws.cell(row=_righe_nota, column=C_LEG,
            value="Le celle grigio scuro sono turni che quel giorno non esistono: "
                  "non vanno compilate.").font = Font(name=CAL, size=10, italic=True)

ws.freeze_panes = "C4"
ws.protection.sheet = False
wb.save(NOME_FILE)

# ---------------- pulizia di una parte che fa "riparare" il file --------------
# openpyxl scrive dentro il file una sezione accessoria (docProps/custom.xml)
# lasciata vuota. Excel la considera un difetto e all'apertura avvisa che il file
# è danneggiato e lo ripara. Non tocca le formule, ma spaventa e basta.
# Qui la si toglie insieme ai due riferimenti che la richiamano.
import zipfile, os, re
def _togli_custom_props(percorso):
    tmp = percorso + ".tmp"
    with zipfile.ZipFile(percorso) as z:
        dati = {n: z.read(n) for n in z.namelist()}
    tolto = False
    # 1) sezione accessoria vuota: fa comparire l'avviso di riparazione
    if "docProps/custom.xml" in dati:
        del dati["docProps/custom.xml"]
        dati["_rels/.rels"] = re.sub(
            r'<Relationship[^>]*custom-properties[^>]*/>', '',
            dati["_rels/.rels"].decode()).encode()
        dati["[Content_Types].xml"] = re.sub(
            r'<Override[^>]*docProps/custom\.xml[^>]*/>', '',
            dati["[Content_Types].xml"].decode()).encode()
        tolto = True
    # 2) i risultati del file VUOTO, scritti accanto alle formule.
    #    openpyxl lascia lì un valore vuoto, ed Excel lo mostra invece di
    #    calcolare: la cella appare bianca. Qui ci si mette il valore giusto,
    #    che per un file senza codici e senza assegnazioni è noto in partenza.
    #    Appena scrivi qualcosa, Excel ricalcola da sé.
    testo = dati["xl/worksheets/sheet1.xml"].decode()

    def valore(rif, v):
        """sostituisce il <v></v> vuoto della cella con il risultato"""
        nonlocal testo
        if isinstance(v, str):
            nuovo = ('<v>' + v.replace("&", "&amp;").replace("<", "&lt;")
                     .replace(">", "&gt;") + '</v>')
            testo = re.sub(r'(<c r="' + rif + r'"[^>]*?)(\s+t="[^"]*")?(>.*?<f[^>]*>.*?</f>)<v></v>',
                           lambda m: m.group(1) + ' t="str"' + m.group(3) + nuovo,
                           testo, count=1, flags=re.S)
        else:
            testo = re.sub(r'(<c r="' + rif + r'"[^>]*>.*?<f[^>]*>.*?</f>)<v></v>',
                           lambda m: m.group(1) + '<v>' + str(v) + '</v>',
                           testo, count=1, flags=re.S)

    tutti_i_nomi = ", ".join(PERSONE)
    for i in range(NG):
        r = R0 + i
        g = calendar.weekday(ANNO, NMESE, i + 1)
        valore("A%d" % r, GG[g])
        valore("%s%d" % (get_column_letter(C_GG), r), GG[g])
        valore("%s%d" % (get_column_letter(C_DD), r), i + 1)
        valore("%s%d" % (get_column_letter(C_DISP), r), tutti_i_nomi)
        valore("%s%d" % (get_column_letter(C_LIB), r), tutti_i_nomi)
        for i2, p in enumerate(PERSONE):
            valore("%s%d" % (get_column_letter(A_D + i2), r), p + ", ")
            valore("%s%d" % (get_column_letter(A_L + i2), r), p + ", ")
    for i2, p in enumerate(PERSONE):
        r = 4 + i2
        for c0 in (CC0, PC0):
            valore("%s%d" % (get_column_letter(c0), r), p)
        for c in range(CC0 + 1, CC1 + 1):
            valore("%s%d" % (get_column_letter(c), r), 0)
        valore("%s%d" % (get_column_letter(A_D + i2), 3), p)
        valore("%s%d" % (get_column_letter(A_L + i2), 3), p)
    for c in list(range(CC0 + 1, CC1 + 1)) + list(range(PC0 + 1, PC1 + 1)):
        valore("%s%d" % (get_column_letter(c), RT), 0)
    # nelle celle il cui risultato è vuoto — CONTROLLO quando non c'è nessuna
    # sovrapposizione — non si lascia un valore memorizzato vuoto: senza, Excel
    # calcola e ottiene comunque vuoto, che è il risultato giusto.
    testo = testo.replace("<v></v>", "").replace("<v/>", "")
    dati["xl/worksheets/sheet1.xml"] = testo.encode()
    # 3) il modo di calcolo: automatico, e NIENTE ricalcolo forzato all'apertura.
    #    Con fullCalcOnLoad Excel butta via i valori memorizzati per ricalcolare;
    #    se il calcolo è impostato su manuale non ricalcola, e le celle restano
    #    vuote. Meglio lasciargli i valori che ci sono e chiedere il calcolo
    #    automatico, che vale poi per le modifiche successive.
    wbx = dati["xl/workbook.xml"].decode()
    wbx = re.sub(r"<calcPr[^>]*/>", '<calcPr calcId="191029" calcMode="auto" '
                 'fullCalcOnLoad="0" forceFullCalc="0"/>', wbx)
    if "<calcPr" not in wbx:
        wbx = wbx.replace("</workbook>", '<calcPr calcId="191029" calcMode="auto"/>'
                          "</workbook>")
    dati["xl/workbook.xml"] = wbx.encode()
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in dati.items():
            z.writestr(n, d)
    os.replace(tmp, percorso)
    return tolto
_pulito = _togli_custom_props(NOME_FILE)

# ------------------------- VERIFICHE ---------------------------------
# Controllano il TESTO delle formule nel file salvato. Non scrivono valori
# nel file: i risultati li calcola Excel quando lo apri.

v = load_workbook(NOME_FILE)          # rilettura del file salvato
w = v[ws.title]
form = [(c.coordinate, c.value) for row in w.iter_rows()
        for c in row if isinstance(c.value, str) and c.value.startswith("=")]
def esito(t, ok, extra=""):
    print(("  OK   " if ok else "  ERRORE ") + t + ((" — " + extra) if extra else ""))

print("=== VERIFICHE SUL FILE SALVATO ===")
print("  file:", NOME_FILE)
print("  OK   risultati del file vuoto scritti accanto alle formule: si vedono")
print("       anche se Excel è impostato su calcolo manuale")
esito(f"righe del mese: {NG}", w.cell(row=R1, column=2).value == NG)
esito("giorno della settimana calcolato con formula",
      str(w["A4"].value).startswith("=CHOOSE"))
esito("nessuna formula contiene TEXTJOIN",
      all("TEXTJOIN" not in f.upper() for _, f in form))
esito("nessuna formula supera 8.192 caratteri",
      all(len(f) <= 8192 for _, f in form),
      f"la più lunga è di {max(len(f) for _, f in form)} caratteri")
esito("nessuna formula ha parentesi attaccate ')('",
      not any(")(" in f for _, f in form),
      "; ".join(rif for rif, f in form if ")(" in f)[:120])
esito("parentesi bilanciate in tutte le formule",
      all(f.count("(") == f.count(")") for _, f in form))
rd = get_column_letter(A_D); rl = get_column_letter(A_L)
esito("i riferimenti di riga cambiano riga per riga",
      f"{P0 and get_column_letter(P0)}5" in str(w[f"{rd}5"].value)
      and f"{get_column_letter(P0)}{R1}" in str(w[f"{rd}{R1}"].value))
esito(f"formula area DISPONIBILI: {len(str(w[rd+'5'].value))} caratteri", True)
esito(f"formula area LIBERI: {len(str(w[rl+'5'].value))} caratteri", True)
esito(f"formula colonna finale: {len(str(w[get_column_letter(C_DISP)+'5'].value))} caratteri", True)
esito("blocco TURNI MESE PRECEDENTE senza formule",
      all(w.cell(row=r, column=c).value is None
          for r in range(4, 4 + len(PERSONE))
          for c in range(PC0 + 1, PC1 + 1)))
esito("celle dei turni vuote", all(w.cell(row=r, column=c).value is None
      for r in range(R0, R1 + 1) for c in range(T0, T1 + 1)))
esito("foglio non protetto", not w.protection.sheet)
esito(f"codici in tendina: {', '.join(CODICI)}",
      len(list(w.data_validations.dataValidation)) == 1)
d0 = list(w.data_validations.dataValidation)[0] if list(w.data_validations.dataValidation) else None
esito(f"messaggio di aiuto entro i 255 caratteri di Excel: {len(d0.prompt or '')}",
      d0 is not None and len(d0.prompt or "") <= 255)
esito(f"messaggio di errore entro i 255 caratteri: {len(d0.error or '')}",
      d0 is not None and len(d0.error or "") <= 255)
esito("nessun testo a capo",
      not any(c.alignment.wrap_text for row in w.iter_rows() for c in row))

print()
print("Posizioni:  turni", f"{LT0}:{LT1}", "| controlli",
      f"{get_column_letter(C_CTRL)}:{get_column_letter(C_LIB)}",
      "| disponibilità", f"{get_column_letter(C_GG)}:{LP1}",
      "| legenda", get_column_letter(C_LEG),
      "| conteggi", f"{get_column_letter(CC0)}:{get_column_letter(PC1)}",
      "| aree di calcolo", f"{rd}:{get_column_letter(ULT)} (nascoste)")
print()
print("APRI IL FILE E CONTROLLA QUESTO: con l'orario vuoto, in ogni riga")
print("DISPONIBILI e LIBERI devono elencare tutti i", len(PERSONE), "nomi.")
print("È la prova che le formule funzionano davvero.")

try:
    from google.colab import files
    files.download(NOME_FILE)
except Exception:
    print("\n(fuori da Colab: il file è stato salvato come", NOME_FILE + ")")
